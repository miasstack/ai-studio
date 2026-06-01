from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUT = "Kongo, K6WYA, John Gotit & CantRushTheVibe tp Montana 700_ TBD_producer_agreement.docx"

ARTIST = "Tony McDowell pka Montana 700"
ARTIST_DISPLAY = "Montana 700"
COMPANY = "Tony McDowell pka Montana 700"
DISTRIBUTOR = "Double M Records LLC Records"
ALBUM = "TBD"
TRACK = "TBD"
AGREEMENT_DATE = "May 5, 2026"
ARTIST_AGREEMENT_DATE = "May 13, 2024"

PRODUCERS = [
    {
        "display": "Kongo",
        "legal": "Joshua Cilumba",
        "address": "",
        "email": "",
        "master": "0.75%",
        "advance": "$625.00",
        "publishing": "10%",
    },
    {
        "display": "K6WYA",
        "legal": "Evan Huang",
        "address": "6635 W 147th Ter\nOverland Park, KS",
        "email": "sohekei@gmail.com",
        "ipi": "1072489249",
        "pro": "BMI",
        "master": "0.75%",
        "advance": "$625.00",
        "publishing": "13.3%",
    },
    {
        "display": "John Gotit",
        "legal": "John Ramirez",
        "address": "6113 Gulf Freeway\nHouston, TX",
        "email": "Fundamentallysound8@gmail.com",
        "master": "0.75%",
        "advance": "$625.00",
        "publishing": "13.3%",
    },
    {
        "display": "CantRushTheVibe",
        "legal": "Randy Jamal Razz",
        "address": "6301 Almeda Rd\nHouston, TX 77004",
        "email": "Fundamentallysound8@gmail.com",
        "master": "0.75%",
        "advance": "$625.00",
        "publishing": "13.3%",
    },
]


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Arial"
    styles["Title"].font.size = Pt(14)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(11)
    styles["Heading 2"].font.bold = True


def p(doc, text="", style=None, bold=False, align=None):
    para = doc.add_paragraph(style=style)
    if align:
        para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    return para


def add_page_break(doc):
    doc.add_page_break()


def producer_ref(producer, include_contact=False):
    base = f"{producer['legal']} p/k/a {producer['display']}"
    if include_contact:
        bits = [base]
        if producer.get("address"):
            bits.append(producer["address"].replace("\n", ", "))
        if producer.get("email"):
            bits.append(producer["email"])
        return " ".join(bits)
    return base


def add_exhibit_heading(doc, label, title):
    p(doc, f'EXHIBIT "{label}"', style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, title, style="Heading 2", align=WD_ALIGN_PARAGRAPH.CENTER)


def add_main_agreement(doc):
    p(doc, "PRODUCER AGREEMENT", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    intro_names = ", ".join(producer_ref(x, include_contact=True) for x in PRODUCERS[:-1])
    intro_names += f" and {producer_ref(PRODUCERS[-1], include_contact=True)}"
    p(
        doc,
        f"{intro_names} (herein collectively referred to as \"Producer\" or \"you\") hereby declares and certifies to {COMPANY} (\"Company\") that Producer will produce and create one (1) original master recording (\"Master\") of the composition(s) contained on Exhibit A for possible inclusion on the next album or other record titled \"{ALBUM}\" (the \"Album\") embodying the featured performance of {ARTIST} (\"Artist\"). In consideration of the advance and the royalty participation set forth below (the \"Royalty Participation\"), and for the express and direct benefit of Company, Artist and {DISTRIBUTOR} (\"Distributor\") and their respective divisions, affiliates, licensees, parents and assigns, Producer hereby declares and certifies as follows:",
    )

    clauses = [
        ("A.", "Producer grants to Company and its designees the perpetual, non-exclusive right to use and publish and to permit others to use and publish Producer's professional names, approved likeness, and approved biographical material concerning Producer for advertising and trade purposes solely in connection with the sale and exploitation of the Master and records manufactured from the Master, or to refrain therefrom."),
        ("B.", "Producer acknowledges and agrees that the Master (excluding the underlying musical composition) embodying the results and proceeds of Producer's services shall, from the inception of creation, constitute a \"work made for hire\" for Company within the meaning of the United States Copyright Act of 1976, as amended. If for any reason the Master does not constitute a work made for hire, then Producer hereby irrevocably transfers and assigns to Company all of Producer's right, title and interest in and to such Master (excluding the underlying musical composition), together with all rights therein. Without limiting the generality of the foregoing, Company and its designees (including, without limitation, Distributor) shall have the exclusive, unrestricted, worldwide and perpetual right (but not the obligation) to use, distribute, sell and exploit the Master in any and all media now known or hereafter invented."),
        ("C.", "Producer and Company warrant and represent to each other that their individual contributions to the Master, including all music, performances and other material (including, without limitation, any so-called \"samples\") furnished in connection with the Master are or will be original with Producer and/or Company, or in the public domain throughout the world or used with the consent of the original owner thereof, and shall not infringe upon or violate any copyright of, or infringe upon or violate the right of privacy or any other right of, any person. Each party agrees to hold the other and its successors, licensees and assigns harmless from and against any and all third party damages, losses, costs and expenses (including reasonable attorney's fees and costs) which the other party or any of its successors, licensees or assigns may suffer or incur by reason of the breach of any of the warranties made herein."),
        ("D. (1)", "In consideration of the rights granted to Company and all services rendered by Producer in connection with the Master, and conditioned upon Producer's performance of all of the material terms and conditions of this Agreement, Company shall instruct Distributor to pay royalties to Producer pursuant to irrevocable letters of direction (each, an \"LOD\"), copies of which are attached hereto, at the royalty rates set forth on Exhibit A."),
        ("D. (2)", "The Producer Royalty shall be calculated, determined, adjusted, reduced, paid and not paid pursuant to the same royalty provisions applicable to Company under its applicable agreement with Distributor. No royalty shall be payable to Producer hereunder until Distributor has recouped all Recording Costs in connection with the Master from the net artist royalties calculated pursuant to such royalty provisions, and the applicable advance shall be recoupable from Producer's royalties."),
        ("D. (3)", "As to records not consisting entirely of the Master, the royalty rate otherwise payable to Producer hereunder with respect to sales of any such record shall be prorated by multiplying such royalty rate by a fraction, the numerator of which is equal to the number of Masters embodied on such record and the denominator of which is the total number of royalty bearing master recordings embodied thereon."),
        ("D. (4)", "Producer acknowledges and agrees that the advance and royalty participation set forth on Exhibit A will be the sole remuneration in respect of the services rendered by Producer in connection with the Master and the exploitation of the Master, and such constitutes fair and equitable remuneration for the services rendered and the rights granted herein."),
        ("E.", "Company shall accord Producer a customary courtesy credit on the liner notes of any Album in all configurations and anywhere else that such credits appear, including metadata, in substantially the form contained on Exhibit A."),
        ("F.", "For a period of five (5) years following release of the Master, Producer agrees that Producer shall not remix, produce and/or co-produce any recording for any person, firm or corporation other than Company which embodies, in whole or in part, any of the selections recorded in the Master."),
        ("G.", "It is hereby agreed that Producer shall be deemed a writer as contained on Exhibit A. Producer shall, and shall cause Producer's respective licensees, designees and affiliates to, grant licenses to Company and Company's respective licensees, designees, affiliates, distributors and assigns in the United States and Canada at a rate equal to one hundred percent (100%) of the minimum statutory copyright royalty rate, determined as of the date the applicable Master is commercially released, for the right to reproduce, use or otherwise exploit the Controlled Compositions in recordings featuring Artist, in any and all configurations in any and all media throughout the universe, subject to any applicable mechanical royalty caps."),
        ("H.", "Producer agrees not to disseminate any copies of the Master, including without limitation reference copies of the Master or copies in any other medium, including Pro Tools files, to any person, except at Company's direction."),
        ("I.", "Upon execution, this agreement shall supersede all prior oral or written understandings between the parties concerning the subject matter hereof. It is intended that this agreement will be replaced with a more formal document, but unless and until such time, if any, this agreement shall represent a complete, binding and fully enforceable contract."),
        ("J.", "The failure by either party to perform any of its obligations hereunder shall not be deemed a breach of this agreement unless such party gives the other party written notice of such failure to perform and such failure is not corrected within thirty (30) days from and after receipt of such notice. In the event of any breach of this agreement by Company, Producer's sole remedy shall be an action at law for damages actually incurred, if any, and in no event shall Producer be entitled to seek equitable or other injunctive relief. Producer is acting as an independent contractor and nothing herein shall constitute a partnership, joint venture, agency or employment relationship."),
        ("K.", "This agreement shall be governed by and construed under the laws and judicial decisions of the State of New York. All claims, disputes or disagreements which may arise out of the interpretation, performance or breach of this agreement shall be submitted exclusively to the jurisdiction of the state courts of the State of New York or the Federal District courts located in New York."),
        ("L.", "This agreement may be executed in any number of counterparts, each of which shall be deemed an original, but all of which shall constitute one document. Delivery of an executed counterpart of a signature page to this agreement by facsimile or pdf shall be effective as delivery of a manually executed counterpart of this agreement."),
        ("M.", "Company will compute royalties hereunder and account to Producer as of each June 30 and December 31 for the prior six (6) months, in respect of each such six (6) month period in which there are sales or returns of records or other exploitations of the Master on which royalties are payable to Producer, on or before the next September 30 with respect to the period ending June 30, and on or before March 31 with respect to the period ending December 31."),
        ("N.", "All royalty statements and other accountings rendered to Producer by or on behalf of Company or Artist shall be conclusively binding upon Producer and not subject to any objection unless specific objection in writing, stating the basis thereof, is given to Artist within twenty-four (24) months from the date such statement or accounting is rendered. Producer shall be foreclosed from maintaining any action, claim or proceeding with respect to any statement or accounting due hereunder unless commenced within thirty-six (36) months after the date such statement or accounting is rendered."),
    ]
    for label, body in clauses:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Inches(-0.35)
        para.paragraph_format.left_indent = Inches(0.35)
        run = para.add_run(label + "\t")
        run.bold = True
        para.add_run(body)

    p(doc, "AGREED:", bold=True)
    for producer in PRODUCERS:
        p(doc, "_____________________________")
        p(doc, producer_ref(producer))
        p(doc, "Dated: _______________________")
        p(doc)
    p(doc, ARTIST)
    p(doc, "_____________________________")
    p(doc, "By: An Authorized Representative")
    p(doc, "Dated: _______________________")


def add_exhibit_a(doc):
    add_page_break(doc)
    add_exhibit_heading(doc, "A", "Producer Royalty, Advance, Credit and Publishing Split")
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Song Title", "Producer", "Legal Name", "Producer Royalty", "Advance", "Publishing Split"]
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = h
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for producer in PRODUCERS:
        row = table.add_row().cells
        vals = [TRACK, producer["display"], producer["legal"], producer["master"], producer["advance"], producer["publishing"]]
        for cell, val in zip(row, vals):
            cell.text = val
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row = table.add_row().cells
    vals = [TRACK, ARTIST_DISPLAY, "Tony McDowell", "N/A", "N/A", "50.1%"]
    for cell, val in zip(row, vals):
        cell.text = val
    p(doc, "*** Publishing split is based on the information provided and totals 100%.")
    p(doc, "Total advance: $2,500.00, allocated as $625.00 per producer.")


def add_regular_lod(doc, label, producer):
    add_page_break(doc)
    add_exhibit_heading(doc, label, f"Letter of Direction - {producer['display']}")
    p(doc, "LETTER OF DIRECTION", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, AGREEMENT_DATE)
    p(doc, "3000 Marcus Ave, Suite 1W5\nLake Success, NY 11042")
    p(doc, "Gentlepersons:")
    p(
        doc,
        f"By the terms and conditions of that certain agreement dated as of {AGREEMENT_DATE} (\"Producer Agreement\"), we have engaged {producer_ref(producer, include_contact=True)} (hereinafter referred to as \"Producer\") to render producing services in connection with the master recording entitled \"{TRACK}\" (the \"Master\") to be delivered under the exclusive recording agreement between you and us dated {ARTIST_AGREEMENT_DATE} (the \"Agreement\"). We hereby request and authorize you, solely as an accommodation to us and Artist, to make payments to, account to, and credit Producer, pursuant to the applicable terms and conditions of the Producer Agreement to which this letter of direction is attached.",
    )
    p(doc, f"All payments to Producer under the Producer Agreement shall be payable and remitted to: {producer_ref(producer, include_contact=True)} or as Producer otherwise instructs you in writing.")
    p(doc, "Your compliance with this authorization will constitute an accommodation to us and Artist alone, and nothing herein shall vest in Producer rights as a beneficiary of or party to the Agreement, this instrument or any other agreement between us. All payments hereunder will constitute payment to us and Artist; you will have no liability by reason of any erroneous payment you may make or failure to comply with this authorization. We will indemnify and hold you harmless against any claims asserted against you and any damages, losses or expenses incurred by you by reason of any such payment or otherwise in connection herewith.")
    p(doc, "Sincerely,")
    p(doc, ARTIST)
    p(doc, "By: _______________________")
    p(doc, "An Authorized Signatory")
    p(doc)
    p(doc, "Producer Acknowledged and Agreed:")
    p(doc, "_____________________________")
    p(doc, producer_ref(producer))
    p(doc, "Dated: _______________________")


SOUNDEXCHANGE_TERMS = [
    "Performer represents and warrants that Performer is the featured recording artist who performed on the sound recording(s) identified on the LOD Repertoire Chart attached hereto as Schedule 1.",
    "Performer represents and warrants that Payee is an individual credited or recognized publicly for the commercially released sound recording identified on the LOD Repertoire Chart.",
    "Performer requests and authorizes SoundExchange to pay to and in the name of Payee an amount equal to the Percentage of the royalties otherwise payable by SoundExchange to Performer in respect of the Recordings, thereby reducing the payments from SoundExchange to Performer.",
    "All monies becoming payable under this Letter of Direction shall be remitted to Payee at the address identified above or as Payee otherwise directs SoundExchange in writing.",
    "SoundExchange will honor a written revocation by Performer of the designation made by this Letter of Direction.",
    "SoundExchange may discontinue making payments under this Letter of Direction at any time, including if checks mailed to Payee's last known address are returned, Performer ceases to be a registrant of SoundExchange, or SoundExchange modifies its policies concerning letters of direction.",
    "Performer acknowledges that SoundExchange is providing payments to Payee solely as an accommodation to Performer but that all royalties distributed by SoundExchange to Payee are taxable to Payee.",
    "SoundExchange may rely conclusively, and shall have no liability when acting, upon any written notice, instruction, other document or signature that is reasonably believed by SoundExchange to be genuine and authorized by Performer.",
    "This Letter of Direction shall be governed by and construed in accordance with the substantive laws of the District of Columbia.",
]


def add_soundexchange_lod(doc, label, producer):
    add_page_break(doc)
    add_exhibit_heading(doc, label, f"SoundExchange Letter of Direction - {producer['display']}")
    p(doc, "SoundExchange, Inc.", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, "Letter of Direction", style="Heading 2", align=WD_ALIGN_PARAGRAPH.CENTER)
    p(doc, "Solely as a service and accommodation to those featured artists entitled to royalties under 17 U.S.C. Section 114(g)(2)(D) who specifically authorize SoundExchange to collect and distribute royalties on their behalf, SoundExchange permits such featured artists to designate that a percentage of the royalties due them from SoundExchange relating to certain sound recordings be remitted to creative personnel credited or recognized publicly for the commercially released sound recording on which the featured artist performs.")
    p(doc, f"*Name of Solo Artist(s) or Group on recording(s): {ARTIST_DISPLAY}")
    p(doc, "*Legal Name of Performer(s) for this LOD: Tony McDowell")
    p(doc, f"*Name of Producer, Mixer or Engineer (\"Payee\") - Only Include One Payment Name: {producer['display']}")
    p(doc, f"*Payee Legal Name: {producer['legal']}")
    p(doc, f"*Payee Address - include c/o's here: {producer.get('address') or 'TBD'}")
    p(doc, f"Payee Telephone Number: TBD")
    p(doc, f"*Payee E-Mail: {producer.get('email') or 'TBD'}")
    p(doc, "*Effective Date: [x] This LOD applies retroactively for all tracks listed on the LOD Repertoire Chart.")
    p(doc, "*Payment Percentage: [x] Percentage varies by each track covered by this LOD. Enter percentages on the LOD Repertoire Chart.")
    p(doc, "TERMS AND CONDITIONS", style="Heading 1")
    for term in SOUNDEXCHANGE_TERMS:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(term)
    p(doc, "ACKNOWLEDGED AND ACCEPTED BY:", style="Heading 2")
    p(doc, "Performer Signature: ______________________________")
    p(doc, "*Performer Printed Legal Name: Tony McDowell")
    p(doc, "Date of Signature: ______________________________")
    p(doc, "OR, SoundExchange Authorized LOD Signatory:")
    p(doc, "SoundExchange Authorized LOD Signatory Printed Name: ______________________________")
    p(doc, "Return the original of this form to:\nSoundExchange, Inc.\n733 10th Street NW, 10th Floor\nWashington, DC 20001\nEmail: accounts@soundexchange.com\nFax: 202.640.5859")
    p(doc, "Schedule 1: Repertoire Chart for Featured Artist Letter of Direction", style="Heading 2")
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Track", "Sound Recording Track Name(s)", "Percentage Share", "Effective Date", "Track Version", "ISRC", "Album/Release"]
    for idx, h in enumerate(headers):
        table.rows[0].cells[idx].text = h
    row = table.add_row().cells
    vals = ["Track 1", TRACK, producer["master"], "", "", "", ALBUM]
    for cell, val in zip(row, vals):
        cell.text = val


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    setup_styles(doc)
    add_main_agreement(doc)
    add_exhibit_a(doc)
    labels = iter(["B", "C", "D", "E", "F", "G", "H", "I"])
    for producer in PRODUCERS:
        add_regular_lod(doc, next(labels), producer)
    for producer in PRODUCERS:
        add_soundexchange_lod(doc, next(labels), producer)
    doc.save(OUT)


if __name__ == "__main__":
    main()
