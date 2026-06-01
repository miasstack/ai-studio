from copy import deepcopy

from docx import Document


TEMPLATE = "Kongo, K6WYA, John Gotit & CantRushTheVibe tp Montana 700_ how I came in_template_format_v2.docx"
OUT = "Kongo, K6WYA, John Gotit & CantRushTheVibe tp Montana 700_ TBD_template_match.docx"

TRACK_OLD = "how I came in"
TRACK_NEW = "TBD"
DATE_OLD = "November 12, 2025"
DATE_NEW = "May 5, 2026"


def set_para_text_keep_format(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_in_paragraph(paragraph, old, new):
    if old not in paragraph.text:
        return
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    # Some Word exports split text oddly. If a simple run-level replacement
    # did not catch it, update the paragraph text while keeping paragraph style.
    if old in paragraph.text:
        set_para_text_keep_format(paragraph, paragraph.text.replace(old, new))


def replace_everywhere(doc, old, new):
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, old, new)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_cell_paragraphs(cell, values):
    # Preserve the existing paragraph/run formatting in this template cell.
    while len(cell.paragraphs) < len(values):
        cell.add_paragraph()
    for idx, value in enumerate(values):
        set_para_text_keep_format(cell.paragraphs[idx], value)
    for idx in range(len(values), len(cell.paragraphs)):
        set_para_text_keep_format(cell.paragraphs[idx], "")


def main():
    doc = Document(TEMPLATE)

    replace_everywhere(doc, TRACK_OLD, TRACK_NEW)
    replace_everywhere(doc, DATE_OLD, DATE_NEW)

    # Exhibit A: keep the exact table but update the business terms.
    exhibit = doc.tables[0]
    set_cell_paragraphs(exhibit.rows[1].cells[0], ['"TBD" '])
    set_cell_paragraphs(exhibit.rows[1].cells[2], ["0.75%", "0.75%", "0.75%", "0.75%", ""])
    set_cell_paragraphs(exhibit.rows[1].cells[3], ["0.75%", "0.75%", "0.75%", "0.75%", ""])
    set_cell_paragraphs(exhibit.rows[1].cells[4], ["$ 625.00", "$ 625.00", "$ 625.00", "$ 625.00", ""])
    set_cell_paragraphs(
        exhibit.rows[1].cells[5],
        [
            "10%",
            "Joshua Cilumba",
            "",
            "13.3%",
            "Evan Huang",
            "",
            "13.3%",
            "John Ramirez",
            "",
            "13.3%",
            "Randy Jamal Razz",
            "",
            "50.1%",
            "Tony McDowell",
            "",
            "",
        ],
    )

    # Four SoundExchange LOD repertoire charts.
    for table in doc.tables[1:5]:
        set_para_text_keep_format(table.rows[5].cells[1].paragraphs[0], '"TBD"')
        set_para_text_keep_format(table.rows[5].cells[2].paragraphs[0], "0.75%")

    # The source template contained stray duplicate signature blocks using the
    # old misspelled Konga name. Remove only those blocks from the main
    # agreement signature section, preserving the surrounding paragraph styles.
    duplicate_blocks = [
        (47, 50),
        (56, 59),
        (65, 68),
        (74, 77),
    ]
    for start, end in reversed(duplicate_blocks):
        for idx in range(end, start - 1, -1):
            delete_paragraph(doc.paragraphs[idx])

    doc.save(OUT)


if __name__ == "__main__":
    main()
